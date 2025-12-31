"""
Excel to CSV Converter

Converts Excel files (.xlsx) to CSV format.
Handles multiple sheets by creating separate CSV files.

Functions:
    convert_file(xlsx_path: Path) -> List[Path]
    convert_all_raw() -> Dict[str, List[Path]]
"""

from __future__ import annotations

import csv
import re
import time
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import openpyxl
import pandas as pd

from shared.db import field_mappings as fm

from .utils import (
    MANIFEST_PATH,
    compute_file_hash,
    ensure_dir,
    filter_superseded_xlsx,
    find_split_parts,
    load_manifest,
    save_manifest,
    split_base_from_stem,
    upsert_file_entry,
    update_manifest_summary,
    utc_now,
)

PIPELINE_ROOT = Path(__file__).resolve().parent.parent
CSV_DIR = PIPELINE_ROOT / "csv"
RAW_DIR = PIPELINE_ROOT.parent / "raw_data"

LARGE_XLSX_BYTES = 20 * 1024 * 1024


def _normalize_header_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    return re.sub(r"\s+", " ", text).lower()


def _build_known_headers() -> Set[str]:
    known: Set[str] = set()
    for mappings in fm.TABLE_MAPPING_REGISTRY.values():
        for mapping in mappings:
            for col in mapping.source_columns:
                known.add(_normalize_header_cell(col))
    return known


def _detect_header_row(ws, known_headers: Set[str], max_rows: int = 50) -> Optional[int]:
    best_row = None
    best_score = -1
    best_match_count = 0
    best_nonempty = 0

    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True), start=1):
        normalized = [_normalize_header_cell(v) for v in row]
        nonempty = [v for v in normalized if v]
        if not nonempty:
            continue
        match_count = sum(1 for v in nonempty if v in known_headers)
        score = (match_count * 10) + len(nonempty)
        if score > best_score:
            best_score = score
            best_row = idx
            best_match_count = match_count
            best_nonempty = len(nonempty)

    if best_row is None:
        return None
    if best_match_count > 0:
        return best_row
    if best_nonempty < 3:
        return None
    if best_match_count == 0 and best_nonempty < 6:
        return None
    return best_row


def _should_skip_sheet(sheet_name: str) -> bool:
    lower = sheet_name.strip().lower()
    return lower in {"index", "indexes", "idx", "summary"}

def _extract_header(ws, header_row: int) -> Tuple[List[str], List[int]]:
    header_cells = next(
        ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True),
        None,
    )
    if not header_cells:
        return [], []
    header = []
    for cell in header_cells:
        header.append("" if cell is None else str(cell).strip())
    while header and header[-1] == "":
        header.pop()
    keep_indexes = [i for i, name in enumerate(header) if name and not name.startswith("Unnamed")]
    filtered_header = [header[i] for i in keep_indexes]
    return filtered_header, keep_indexes


def _write_sheet_csv(ws, out_path: Path, header_row: int) -> bool:
    header, keep_indexes = _extract_header(ws, header_row)
    if not header:
        return False
    with out_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            values = []
            has_value = False
            for idx in keep_indexes:
                cell = row[idx] if idx < len(row) else None
                if cell is None:
                    values.append("")
                else:
                    cell_text = str(cell)
                    if cell_text.strip():
                        has_value = True
                    values.append(cell_text)
            if has_value:
                writer.writerow(values)
    return True

def _slugify_sheet(name: str) -> str:
    """Make a sheet name safe for file paths."""
    slug = re.sub(r"[^A-Za-z0-9]+", "_", name.strip())
    return slug.strip("_") or "sheet"


def _slugify_filename(name: str) -> str:
    """Slugify base file names (replace spaces/specials with underscores)."""
    slug = re.sub(r"[^A-Za-z0-9]+", "_", name.strip())
    return slug.strip("_") or "file"


def convert_file(xlsx_path: Path, force: bool = False) -> List[Path]:
    """
    Convert one Excel file to one-or-more CSV files (one per sheet).

    Returns list of CSV paths created.
    """
    start = time.perf_counter()
    if not xlsx_path.exists():
        raise FileNotFoundError(f"File not found: {xlsx_path}")

    if not force:
        split_parts = find_split_parts(xlsx_path)
        if split_parts and split_base_from_stem(xlsx_path.stem) is None:
            print(
                f"[convert] {xlsx_path.name}: skip (split parts detected: "
                f"{len(split_parts)})"
            )
            return []

    manifest = load_manifest(MANIFEST_PATH)
    manifest_index = {entry.get("file_name"): entry for entry in manifest.get("files", [])}
    existing = manifest_index.get(xlsx_path.name)
    if existing and not force:
        existing_hash = existing.get("file_hash")
        outputs = existing.get("csv_outputs") or []
        csv_exist = outputs and all(Path(p).exists() for p in outputs)
        current_hash = compute_file_hash(xlsx_path)
        if csv_exist and existing_hash == current_hash:
            print(f"[convert] {xlsx_path.name}: skip (hash unchanged, CSV present)")
            return [Path(p) for p in outputs if Path(p).exists()]

    ensure_dir(CSV_DIR)
    known_headers = _build_known_headers()
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    use_streaming = xlsx_path.stat().st_size >= LARGE_XLSX_BYTES

    csv_paths: List[Path] = []
    base_name = _slugify_filename(xlsx_path.stem)

    for sheet_name in wb.sheetnames:
        if _should_skip_sheet(sheet_name):
            continue
        ws = wb[sheet_name]
        header_row = _detect_header_row(ws, known_headers)
        if header_row is None:
            continue
        safe_sheet = _slugify_sheet(sheet_name)
        out_path = CSV_DIR / f"{base_name}__{safe_sheet}.csv"
        if use_streaming:
            if _write_sheet_csv(ws, out_path, header_row):
                csv_paths.append(out_path)
            continue
        df = pd.read_excel(xlsx_path, sheet_name=sheet_name, header=header_row - 1, dtype=str)
        if df.empty:
            continue
        df.columns = [str(col).strip() if col is not None else "" for col in df.columns]
        df = df.loc[:, [col for col in df.columns if col and not str(col).startswith("Unnamed")]]
        if df.empty or not df.columns.tolist():
            continue
        df = df.dropna(how="all")
        if df.empty:
            continue
        df.to_csv(out_path, index=False)
        csv_paths.append(out_path)
    wb.close()

    # Update manifest once per source file
    manifest = load_manifest(MANIFEST_PATH)
    entry = upsert_file_entry(manifest, xlsx_path.name)
    entry["status"] = "csv"
    entry["csv_outputs"] = [str(p) for p in csv_paths]
    entry["converted_at"] = utc_now()
    entry["file_hash"] = compute_file_hash(xlsx_path)
    update_manifest_summary(manifest)
    save_manifest(manifest, MANIFEST_PATH)

    elapsed = time.perf_counter() - start
    print(f"[convert] {xlsx_path.name}: {len(csv_paths)} sheets -> CSV in {elapsed:.2f}s")
    return csv_paths


def convert_all_raw() -> Dict[str, List[Path]]:
    """
    Convert all XLSX files under raw_data/ into CSVs.

    Returns dict: {xlsx_name: [csv_paths...]}
    """
    ensure_dir(RAW_DIR)
    manifest = load_manifest(MANIFEST_PATH)
    manifest_index = {entry.get("file_name"): entry for entry in manifest.get("files", [])}

    results: Dict[str, List[Path]] = {}
    xlsx_files, skipped = filter_superseded_xlsx(RAW_DIR.glob("*.xlsx"))
    for skipped_path in skipped:
        print(
            f"[convert-all] skip {skipped_path.name} (split parts detected)"
        )
    for xlsx in xlsx_files:
        entry = manifest_index.get(xlsx.name)
        if entry:
            existing_hash = entry.get("file_hash")
            outputs = entry.get("csv_outputs") or []
            csv_exist = outputs and all(Path(p).exists() for p in outputs)
            current_hash = compute_file_hash(xlsx)
            if csv_exist and existing_hash == current_hash:
                print(f"[convert-all] skip {xlsx.name} (hash unchanged, CSV present)")
                results[xlsx.name] = [Path(p) for p in outputs if Path(p).exists()]
                continue
        csvs = convert_file(xlsx)
        results[xlsx.name] = csvs
    # Refresh summary after conversions/skips
    manifest = load_manifest(MANIFEST_PATH)
    update_manifest_summary(manifest)
    save_manifest(manifest, MANIFEST_PATH)
    return results
