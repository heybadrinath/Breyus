#!/usr/bin/env python3
"""
Split large XLSX files into multiple smaller XLSX parts by row count.

This preserves the original data while keeping each output file under
the target size (approximate). Output files are written next to the input.
"""

from __future__ import annotations

import argparse
import math
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook


def _safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("_")


def _write_part(
    output_path: Path,
    sheet_name: str,
    header: list,
    rows_iter,
    max_rows: int,
) -> int:
    """
    Stream rows into a new XLSX file until max_rows is reached.
    Returns number of data rows written.
    """
    wb_out = Workbook(write_only=True)
    ws_out = wb_out.create_sheet(title=sheet_name)
    if header:
        ws_out.append(header)

    written = 0
    for row in rows_iter:
        ws_out.append(row)
        written += 1
        if written >= max_rows:
            break

    wb_out.save(output_path)
    return written


def split_xlsx(path: Path, target_mb: float) -> None:
    file_size = path.stat().st_size
    target_bytes = int(target_mb * 1024 * 1024)
    if file_size <= target_bytes:
        print(f"SKIP (already under target): {path}")
        return

    wb = load_workbook(path, read_only=True, data_only=True)
    total_rows = 0
    for ws in wb.worksheets:
        total_rows += ws.max_row or 0

    if total_rows <= 1:
        print(f"SKIP (not enough rows): {path}")
        return

    # Estimate rows per part based on file size.
    ratio = target_bytes / max(file_size, 1)
    rows_per_part = max(int(math.floor(total_rows * ratio)), 1)

    print(
        f"SPLIT {path} | size={file_size/1_048_576:.2f}MB "
        f"total_rows={total_rows} rows_per_part~{rows_per_part}"
    )

    for ws in wb.worksheets:
        sheet_name = ws.title or "Sheet1"
        safe_sheet = _safe_name(sheet_name) or "Sheet1"

        rows = ws.iter_rows(values_only=True)
        try:
            header = list(next(rows))
        except StopIteration:
            continue

        part_index = 1
        while True:
            output_name = (
                f"{path.stem}__{safe_sheet}_part{part_index:02d}{path.suffix}"
            )
            output_path = path.with_name(output_name)

            written = _write_part(output_path, sheet_name, header, rows, rows_per_part)
            print(f"  wrote {output_path.name} ({written} rows)")

            if written < rows_per_part:
                break
            part_index += 1


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Split XLSX files into smaller parts by row count."
    )
    parser.add_argument(
        "files",
        nargs="+",
        type=Path,
        help="Paths to XLSX files to split.",
    )
    parser.add_argument(
        "--target-mb",
        type=float,
        default=85.0,
        help="Target max size per part in MB (approx). Default: 85.",
    )

    args = parser.parse_args()

    for file_path in args.files:
        if not file_path.exists():
            print(f"SKIP (missing): {file_path}")
            continue
        split_xlsx(file_path, args.target_mb)


if __name__ == "__main__":
    main()
